"""Onboarding workbook generation tests (single two-sheet Excel Generation).

LOCAL / MOCKED — nothing is uploaded anywhere. These verify the single
onboarding workbook: exactly 13-column ``OB Format`` sheet + exactly 14-column
``Mail Format`` sheet in ONE ``Onboarding_<ts>.xlsx`` file (kind
``excel_generation``), correct values/PII placement (full Aadhaar only inside
the Mail Format sheet), real Excel dates, and clean/transactional failure
handling.

Run with:
    python tests/test_onboarding_pair.py
"""

import datetime
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)) + "/..")

import openpyxl

from app import database as db
from app import generation
from app import master_data
from app import rules

PASS = 0
FAIL = 0


def check(label, cond, extra=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  [PASS] {label}")
    else:
        FAIL += 1
        print(f"  [FAIL] {label} {extra}")


def fresh_db():
    tmp = tempfile.mkdtemp()
    db.DB_DIR = Path(tmp) / "database"
    db.DB_PATH = db.DB_DIR / "teamhr.db"
    db.init_db()
    os.environ["TEAMHR_CONFIG_FILE"] = str(Path(tmp) / "config.json")
    return tmp


def _candidate(batch_id, name, mobile, cost_code, designation, team, operation,
               facility=None, location=None):
    hubs = rules.get_hubs_for_cost_code(cost_code)
    if facility is None:
        facility = hubs[0] if hubs else ""
    if location is None:
        location = master_data.get_location_for_facility(facility) if facility else ""
    return {
        "batch_id": batch_id,
        "name": name,
        "mobile": mobile,
        "designation": designation,
        "facility_name": facility,
        "location_code": location,
        "cost_code": cost_code,
        "status": "ready",
        "entity": "Flipkart" if cost_code in ("4421", "4441") else "Myntra",
        "team": team,
        "operation": operation,
        "migrant": "No",
        "salary": 18000,
        "aadhaar_number": "1234 5678 9012",
        "dob": "15/05/1995",
        "address": "Flat 5, 12th Main Road, Bengaluru",
        "doj": "10/09/2026",
        "gender": "Male",
        "pin_code": "560001",
        "father_name": "",
    }


def SO_PAIR_CANDIDATES(batch_id):
    # Explicit facilities/locations (deterministic, independent of master row
    # ordering): Flipkart LM 4421, Flipkart FM 4441 (Pickup Hub), Myntra LM 8751.
    return [
        _candidate(batch_id, "Anil Kumar", "9000000001", "4421",
                   "LM - Delivery Executive", "LAST MILE - OPERATIONS", "Last Mile",
                   facility="Peenya Hub", location="BLR/PEN"),
        _candidate(batch_id, "Bharath K", "9000000002", "4441",
                   "FM - Delivery Executive", "FIRST MILE - OPERATIONS", "First Mile",
                   facility="NelamangalaHub_BLR_PL", location="NelamangalaHub_BLR_PL"),
        _candidate(batch_id, "Chandan G", "9000000003", "8751",
                   "LM - Delivery Executive", "LAST MILE - OPERATIONS", "Last Mile",
                   facility="BanaswadiMYNTRAHub_BLR", location="BNS/BLR"),
    ]


def setup_case(tmp, candidates):
    gen_dir = Path(tmp) / "gen"
    generation.set_output_base_dir(str(gen_dir))
    db.get_or_create_batch(batch_id=1)
    ids = [db.insert_candidate(c) for c in candidates]
    return gen_dir, ids


def test_generation_shape():
    """Generating lands ONE workbook (both sheets), a single DB row, pair id."""
    tmp = fresh_db()
    gen_dir, ids = setup_case(tmp, SO_PAIR_CANDIDATES(1))

    res = generation.generate_onboarding_workbook(1, only_ready=True)
    check("generation success", res.get("success") is True, str(res))

    path = Path(res["file_path"])
    check("single workbook written", path.exists(), str(path))
    check("workbook named Onboarding_<id>", path.name.startswith("Onboarding_"),
          path.name)
    check("no sibling pair file", len(list(gen_dir.rglob("TeamHR_OB_*.xlsx"))) == 0)

    ts = path.stem.replace("Onboarding_", "")
    check("single pair id PO-<ts>", res["generation_pair_id"] == f"PO-{ts}",
          res["generation_pair_id"])

    gf = db.get_generated_file(res["generated_file_id"])
    check("one kind excel_generation row", gf["kind"] == "excel_generation",
          str(gf.get("kind")))
    check("pair id on row", gf["generation_pair_id"] == f"PO-{ts}",
          str(gf.get("generation_pair_id")))
    check("candidates marked generated",
          db.get_candidate(ids[0])["excel_generated"] == "true")
    check("batch status Generated",
          (db.get_batch(1) or {}).get("status") == "Generated",
          str((db.get_batch(1) or {}).get("status")))

    pairs = db.list_generation_pairs()
    check("list_generation_pairs returns the pair", len(pairs) == 1, str(len(pairs)))
    check("pair id matches", pairs[0]["generation_pair_id"] == f"PO-{ts}")

    check("daily master exported",
          Path(res["daily_master_file"]).exists() if res.get("daily_master_file") else False,
          str(res.get("daily_master_file")))
    return path


def _find_header_row(ws):
    for r in range(1, min(ws.max_row, 6) + 1):
        vals = [str(ws.cell(row=r, column=c).value or "") for c in range(1, ws.max_column + 1)]
        if any(v.strip() for v in vals):
            return r, vals
    return 1, []


def _read_workbook(path):
    wb = openpyxl.load_workbook(str(path))
    ob = wb[generation.OB_FORMAT_SHEET]
    mail = wb[generation.MAIL_FORMAT_SHEET]
    ob_hdr, ob_vals = _find_header_row(ob)
    mail_hdr, mail_vals = _find_header_row(mail)
    return wb, ob, mail, ob_hdr, mail_hdr, ob_vals, mail_vals


def test_ob_sheet_content():
    """OB Format sheet: exact 13 columns + exact per-candidate values."""
    tmp = fresh_db()
    gen_dir, ids = setup_case(tmp, SO_PAIR_CANDIDATES(1))
    res = generation.generate_onboarding_workbook(1, only_ready=True)
    path = Path(res["file_path"])

    wb, ob, mail, ob_hdr, mail_hdr, ob_vals, mail_vals = _read_workbook(path)
    expected_ob = ["Sl No", "Name*", "Mobile Number*", "Team*", "Cost Code*",
                   "Facility Type*", "Line of Business*", "Sub Type*",
                   "Role - Designation*", "Fixed Net Take Home*", "State*",
                   "Facility*", "Contractor*"]
    check("exact 13 OB headers", ob_vals == expected_ob, str(ob_vals))
    check("no extra OB columns", len(ob_vals) == 13, str(len(ob_vals)))

    rows = []
    for r in range(ob_hdr + 1, ob.max_row + 1):
        name = ob.cell(row=r, column=2).value
        if name not in (None, ""):
            rows.append({col: ob.cell(row=r, column=c).value
                         for c, col in enumerate(expected_ob, start=1)})
    check("3 OB data rows", len(rows) == 3, str(len(rows)))
    by_name = {r["Name*"]: r for r in rows}
    check("Anil/Bharath/Chandan present",
          all(n in by_name for n in ("Anil Kumar", "Bharath K", "Chandan G")))

    for nm, mob, cc, team, ft, loc in [
        ("Anil Kumar", "9000000001", "4421", "LAST MILE - OPERATIONS",
         "DELIVERY_HUB", "BLR/PEN"),
        ("Bharath K", "9000000002", "4441", "FIRST MILE - OPERATIONS",
         "PICKUP_HUB", "NelamangalaHub_BLR_PL"),
        ("Chandan G", "9000000003", "8751", "LAST MILE - OPERATIONS",
         "DELIVERY_HUB", "BNS/BLR"),
    ]:
        r = by_name[nm]
        check(f"{nm} mobile @text", str(r["Mobile Number*"]) == mob, str(r["Mobile Number*"]))
        check(f"{nm} cost code", r["Cost Code*"] == cc, str(r["Cost Code*"]))
        check(f"{nm} team", r["Team*"] == team, str(r["Team*"]))
        check(f"{nm} facility type", r["Facility Type*"] == ft, str(r["Facility Type*"]))
        check(f"{nm} LOB", r["Line of Business*"] == "EKART", str(r["Line of Business*"]))
        check(f"{nm} sub type", r["Sub Type*"] == "EKART", str(r["Sub Type*"]))
        check(f"{nm} designation", r["Role - Designation*"] == by_name[nm]["Role - Designation*"])
        check(f"{nm} state", r["State*"] == "KARNATAKA", str(r["State*"]))
        check(f"{nm} facility is location code", r["Facility*"] == loc, str(r["Facility*"]))
        check(f"{nm} contractor",
              r["Contractor*"] == "TEAM HR GSA PRIVATE LIMITED", str(r["Contractor*"]))
        check(f"{nm} salary numeric", r["Fixed Net Take Home*"] == 18000,
              str(r["Fixed Net Take Home*"]))

    # PII: full Aadhaar must NEVER appear inside the OB Format sheet.
    aadhar = "1234 5678 9012"
    leak = []
    for row in ob.iter_rows(values_only=True):
        for cell in row:
            if cell is not None and aadhar in str(cell):
                leak.append(str(cell))
    check("OB sheet never contains full Aadhaar", not leak, str(leak[:1]))


def test_no_aadhaar_in_name_or_folders():
    tmp = fresh_db()
    gen_dir, ids = setup_case(tmp, SO_PAIR_CANDIDATES(1))
    res = generation.generate_onboarding_workbook(1, only_ready=True)
    for piece in (res["filename"], res["file_path"], res["date_folder"]):
        check(f"no Aadhaar in '{piece}'", "1234" not in piece and "9012" not in piece)


def test_mail_sheet_content():
    """Mail Format sheet: exact 14 headers + dates as dates + full PII here."""
    tmp = fresh_db()
    gen_dir, ids = setup_case(tmp, SO_PAIR_CANDIDATES(1))
    res = generation.generate_onboarding_workbook(1, only_ready=True)
    path = Path(res["file_path"])

    wb, ob, mail, ob_hdr, mail_hdr, ob_vals, mail_vals = _read_workbook(path)
    expected_mail = ["Date of Joining", "Name", "Mobile No", "Designation",
                     "Branch", "Vertical", "State", "Net Salary", "Aadhar No",
                     "DOB", "Fathers Name", "Address", "Pin Code", "Gender"]
    check("exact 14 mail headers", mail_vals == expected_mail, str(mail_vals))
    check("no recruiter/uan column", "Recruiter" not in str(mail_vals)
          and "UAN" not in str(mail_vals).upper(), str(mail_vals))

    rows = []
    for r in range(mail_hdr + 1, mail.max_row + 1):
        if mail.cell(row=r, column=2).value not in (None, ""):
            rows.append({col: mail.cell(row=r, column=c).value
                         for c, col in enumerate(expected_mail, start=1)})
    check("3 mail rows", len(rows) == 3, str(len(rows)))
    by_name = {r["Name"]: r for r in rows}
    check("all three present", all(n in by_name for n in ("Anil Kumar", "Bharath K", "Chandan G")))

    anil_row = mail_hdr + 1 + rows.index(by_name["Anil Kumar"])
    r = by_name["Anil Kumar"]
    check("designation", r["Designation"] == "LM - Delivery Executive", str(r["Designation"]))
    check("branch is location code", r["Branch"] == "BLR/PEN", str(r["Branch"]))
    check("vertical is facility name", r["Vertical"] == "Peenya Hub", str(r["Vertical"]))
    check("state title-cased Karnataka", r["State"] == "Karnataka", str(r["State"]))
    check("net salary numeric", r["Net Salary"] == 18000, str(r["Net Salary"]))
    check("aadhaar full present", str(r["Aadhar No"]).replace(" ", "") == "123456789012",
          str(r["Aadhar No"]))
    check("father name blank", r["Fathers Name"] in ("", None), str(r["Fathers Name"]))
    check("address", r["Address"] == "Flat 5, 12th Main Road, Bengaluru", str(r["Address"]))
    check("pin code", str(r["Pin Code"]) == "560001", str(r["Pin Code"]))
    check("gender", r["Gender"] == "Male", str(r["Gender"]))

    # Dates as real Excel dates with DD/MM/YYYY format; columns numbered per spec.
    doj_cell = mail.cell(row=anil_row, column=1)
    dob_cell = mail.cell(row=anil_row, column=10)
    check("DOJ is a date", isinstance(doj_cell.value, (datetime.date, datetime.datetime)),
          str(doj_cell.value))
    check("DOJ format DD/MM/YYYY", doj_cell.number_format == "DD/MM/YYYY",
          doj_cell.number_format)
    check("DOJ value 2026-09-10", str(doj_cell.value)[:10] == "2026-09-10", str(doj_cell.value))
    check("DOB is a date", isinstance(dob_cell.value, (datetime.date, datetime.datetime)),
          str(dob_cell.value))
    check("DOB value 1995-05-15", str(dob_cell.value)[:10] == "1995-05-15", str(dob_cell.value))

    check("mobile TEXT", mail.cell(row=anil_row, column=3).number_format == "@")
    check("aadhar TEXT", mail.cell(row=anil_row, column=9).number_format == "@")
    check("salary #,##0", mail.cell(row=anil_row, column=8).number_format == "#,##0")


def test_generation_failure_clean():
    """A totally invalid candidate (8752, no Myntra FM master) fails cleanly."""
    tmp = fresh_db()
    gen_dir, ids = setup_case(tmp, [
        _candidate(1, "Failing FM", "9000000099", "8752",
                   "FM - Delivery Executive", "FIRST MILE - OPERATIONS", "First Mile",
                   facility="", location=""),
    ])
    res = generation.generate_onboarding_workbook(1, only_ready=True)
    check("generation fails", res.get("success") is False, str(res))
    check("error explains validation",
          "No Ready candidates" in res.get("error", "") or "No Myntra" in res.get("error", ""),
          res.get("error", ""))
    cand = db.get_candidate(ids[0])
    check("candidate moved to needs_attention", cand["status"] == "needs_attention",
          str(cand["status"]))
    files = [p for p in gen_dir.rglob("*") if p.is_file()]
    check("no files left behind", not files, str(files))


def test_generation_transactional_no_partial_files():
    """Failure during build leaves NO file on disk and no DB rows."""
    tmp = fresh_db()
    gen_dir, ids = setup_case(tmp, SO_PAIR_CANDIDATES(1))
    from app import generation as g
    orig = g.build_onboarding_workbook
    def boom(rows, mail_rows):
        raise RuntimeError("simulated workbook build failure")
    g.build_onboarding_workbook = boom
    try:
        res = g.generate_onboarding_workbook(1, only_ready=True)
    finally:
        g.build_onboarding_workbook = orig
    check("generation reports failure", res.get("success") is False, str(res))
    files = [p for p in gen_dir.rglob("*") if p.is_file()]
    check("no partial files remain", not files, str(files))
    pairs = db.list_generation_pairs()
    check("no generation rows recorded", not pairs, str([p["generation_pair_id"] for p in pairs]))


def test_missing_mail_field_blocks():
    """Missing gender/pin/doj blocks the WORKBOOK (no OB-only output)."""
    tmp = fresh_db()
    gen_dir, ids = setup_case(tmp, SO_PAIR_CANDIDATES(1))
    cand = _candidate(1, "Missing Fields", "9000000088", "4421",
                      "LM - Delivery Executive", "LAST MILE - OPERATIONS", "Last Mile",
                      facility="Peenya Hub", location="BLR/PEN")
    for k in ("doj", "gender", "pin_code"):
        cand[k] = ""
    db.insert_candidate(cand)
    res = generation.generate_onboarding_workbook(1, only_ready=True)
    check("generation blocked", res.get("success") is False, str(res))
    check("error cites mail requirements",
          "Mail Format requirements" in res.get("error", ""), res.get("error", ""))
    files = [p for p in gen_dir.rglob("*") if p.is_file()]
    check("no partial files", not files, str(files))


def test_invalid_candidate_excluded():
    """A candidate failing OB validation is excluded and noted, never written."""
    tmp = fresh_db()
    gen_dir, ids = setup_case(tmp, SO_PAIR_CANDIDATES(1))
    cand = _candidate(1, "Bad Designation", "9000000077", "4421",
                      "LM - Not A Real Role", "LAST MILE - OPERATIONS", "Last Mile",
                      facility="Peenya Hub", location="BLR/PEN")
    bad_id = db.insert_candidate(cand)
    res = generation.generate_onboarding_workbook(1, only_ready=True)
    check("generation still succeeds for the valid candidates",
          res.get("success") is True, str(res))
    check("bad candidate marked needs_attention",
          db.get_candidate(bad_id)["status"] == "needs_attention",
          str(db.get_candidate(bad_id)["status"]))
    check("valid candidate still ready", db.get_candidate(ids[0])["status"] == "ready")
    path = Path(res["file_path"])
    wb, ob, mail, ob_hdr, mail_hdr, _, _ = _read_workbook(path)
    data_rows = [r for r in range(ob_hdr + 1, ob.max_row + 1)
                 if ob.cell(row=r, column=2).value not in (None, "")]
    check("only the 3 valid rows written", len(data_rows) == 3, str(len(data_rows)))


def test_repeat_generation_no_overwrite():
    """Two generations create distinct workbooks (never overwrite)."""
    tmp = fresh_db()
    gen_dir, ids = setup_case(tmp, SO_PAIR_CANDIDATES(1))
    res1 = generation.generate_onboarding_workbook(1, only_ready=True)
    res2 = generation.generate_onboarding_workbook(1, only_ready=True)
    check("both generations succeed", res1["success"] and res2["success"])
    check("filenames differ", res1["filename"] != res2["filename"],
          f"{res1['filename']} / {res2['filename']}")
    check("pair ids differ", res1["generation_pair_id"] != res2["generation_pair_id"])
    rows = db.list_generated_files(batch_id=1, limit=10)
    check("two excel_generation rows",
          len(rows) == 2 and all(r["kind"] == "excel_generation" for r in rows),
          str([(r["kind"]) for r in rows]))
    all_wb = list(gen_dir.rglob("Onboarding_*.xlsx"))
    check("two workbooks", len(all_wb) == 2, str([p.name for p in all_wb]))


def test_recruiter_profile():
    """Recruiter profile persists to isolated config and drives the API."""
    tmp = fresh_db()
    generation.save_recruiter_profile("Recruiter Default")
    check("profile saved", generation.get_default_recruiter_name() == "Recruiter Default",
          generation.get_default_recruiter_name())
    generation.save_recruiter_profile("")
    check("profile cleared", generation.get_default_recruiter_name() == "",
          generation.get_default_recruiter_name())
    generation.save_recruiter_profile("Recruiter Default")


def test_daily_master_masks_aadhaar():
    """Daily master export never contains a full Aadhaar."""
    tmp = fresh_db()
    gen_dir, ids = setup_case(tmp, SO_PAIR_CANDIDATES(1))
    res = generation.generate_onboarding_workbook(1, only_ready=True)
    daily = Path(res["daily_master_file"])
    check("daily master exists", daily.exists(), str(daily))
    wb = openpyxl.load_workbook(str(daily))
    ws = wb.worksheets[0]
    full = "123456789012"
    leak = []
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if cell is not None and full in str(cell).replace(" ", ""):
                leak.append(str(cell))
    check("daily master has no full Aadhaar", not leak, str(leak[:1]))


def test_onboarding_api():
    """/api/generate-excel generates; download works; generated-pairs hides PII."""
    from fastapi.testclient import TestClient
    from app import main

    tmp = fresh_db()
    gen_dir, ids = setup_case(tmp, SO_PAIR_CANDIDATES(1))
    client = TestClient(main.app)

    r = client.post("/api/generate-excel", json={"batch_id": 1})
    check("API generation 200", r.status_code == 200, str(r.status_code))
    data = r.json()
    check("API success", data.get("success") is True, str(data))
    check("API returns file id", data.get("generated_file_id"))

    gid = data["generated_file_id"]
    d = client.get(f"/api/generated-file/{gid}/download")
    check("download route 200", d.status_code == 200, str(d.status_code))

    gp = client.get("/api/generated-pairs").json()
    check("generated-pairs returns pair", len(gp.get("pairs", [])) >= 1)
    with_full = False
    for pair in gp["pairs"]:
        if "123456789012" in str(pair):
            with_full = True
    check("generated-pairs never leaks full Aadhaar", not with_full)


def run_all():
    print("=" * 70)
    print("ONBOARDING WORKBOOK: SINGLE TWO-SHEET EXCEL GENERATION")
    print("=" * 70)
    test_generation_shape()
    print("\n--- OB Format sheet content ---")
    test_ob_sheet_content()
    print("\n--- No Aadhaar in names/folders ---")
    test_no_aadhaar_in_name_or_folders()
    print("\n--- Mail Format sheet content ---")
    test_mail_sheet_content()
    print("\n--- Failure handling ---")
    test_generation_failure_clean()
    test_generation_transactional_no_partial_files()
    test_missing_mail_field_blocks()
    test_invalid_candidate_excluded()
    print("\n--- No overwrite ---")
    test_repeat_generation_no_overwrite()
    print("\n--- Recruiter profile ---")
    test_recruiter_profile()
    print("\n--- Daily master masks Aadhaar ---")
    test_daily_master_masks_aadhaar()
    print("\n--- API ---")
    test_onboarding_api()

    print("=" * 70)
    print(f"TOTAL: {PASS + FAIL}  PASS: {PASS}  FAIL: {FAIL}")
    if FAIL:
        sys.exit(1)


if __name__ == "__main__":
    run_all()